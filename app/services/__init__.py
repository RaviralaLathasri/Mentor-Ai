"""
Service layer for AI Mentor backend.

The routes call these classes for business logic. This file intentionally keeps
all current service classes in one module to preserve existing imports.
"""

from __future__ import annotations

import logging
import os
import random
import re
import time
import uuid
from datetime import datetime
from typing import Dict, List, Optional, Tuple

from sqlalchemy import desc, func
from sqlalchemy.orm import Session

from app.database import (
    AdaptiveSession,
    DifficultyLevel,
    Feedback,
    FeedbackType,
    MentorResponse,
    Student,
    StudentProfile,
    WeaknessScore,
)
from app.schemas import (
    AdaptationUpdate,
    MistakeExplanation,
    ResumeIssue,
    ResumeMentorResponse,
    ResumeSectionAnalysis,
    StudentContextSnapshot,
    WeaknessAnalysisResult,
)
from app.services.resume_insights import (
    AI_DATA_KEYWORDS,
    calculate_resume_score,
    improvement_suggestions,
    keyword_gap_analysis,
)
from app.utils.openai_client import get_openai_client

logger = logging.getLogger(__name__)


def _coerce_difficulty(value: Optional[object], default: DifficultyLevel = DifficultyLevel.MEDIUM) -> DifficultyLevel:
    if value is None:
        return default
    if isinstance(value, DifficultyLevel):
        return value
    raw = value.value if hasattr(value, "value") else str(value)
    return DifficultyLevel(raw)


def _coerce_feedback_type(value: object) -> FeedbackType:
    if isinstance(value, FeedbackType):
        return value
    raw = value.value if hasattr(value, "value") else str(value)
    return FeedbackType(raw)


class StudentProfileService:
    """Create/read/update student profiles."""

    def __init__(self, db: Session):
        self.db = db

    def create_profile(
        self,
        student_id: int,
        skills: Optional[List[str]] = None,
        interests: Optional[List[str]] = None,
        goals: str = "",
        confidence_level: float = 0.5,
        preferred_difficulty: object = DifficultyLevel.MEDIUM,
    ) -> StudentProfile:
        student = self.db.query(Student).filter(Student.id == student_id).first()
        if not student:
            raise ValueError(f"Student {student_id} not found")

        existing = self.get_profile(student_id)
        if existing:
            raise ValueError(f"Profile already exists for student {student_id}")

        if not (0.0 <= confidence_level <= 1.0):
            raise ValueError("Confidence must be between 0.0 and 1.0")

        profile = StudentProfile(
            student_id=student_id,
            skills=skills or [],
            interests=interests or [],
            goals=goals or "",
            confidence_level=confidence_level,
            preferred_difficulty=_coerce_difficulty(preferred_difficulty),
        )
        self.db.add(profile)
        self.db.commit()
        self.db.refresh(profile)
        return profile

    def get_profile(self, student_id: int) -> Optional[StudentProfile]:
        return self.db.query(StudentProfile).filter(StudentProfile.student_id == student_id).first()

    def update_profile(
        self,
        student_id: int,
        skills: Optional[List[str]] = None,
        interests: Optional[List[str]] = None,
        goals: Optional[str] = None,
        confidence_level: Optional[float] = None,
        preferred_difficulty: Optional[object] = None,
    ) -> StudentProfile:
        profile = self.get_profile(student_id)
        if not profile:
            raise ValueError(f"Profile not found for student {student_id}")

        if skills is not None:
            profile.skills = skills
        if interests is not None:
            profile.interests = interests
        if goals is not None:
            profile.goals = goals
        if confidence_level is not None:
            if not (0.0 <= confidence_level <= 1.0):
                raise ValueError("Confidence must be between 0.0 and 1.0")
            profile.confidence_level = confidence_level
        if preferred_difficulty is not None:
            profile.preferred_difficulty = _coerce_difficulty(preferred_difficulty)

        profile.updated_at = datetime.utcnow()
        self.db.commit()
        self.db.refresh(profile)
        return profile

    def get_learning_context(self, student_id: int) -> Dict:
        profile = self.get_profile(student_id)
        if not profile:
            return {}
        return profile.learning_style_summary


class WeaknessAnalyzerService:
    """Weakness-first analyzer and explain-my-mistake logic."""

    _QUIZ_QUESTION_BANK: Dict[str, List[Dict[str, object]]] = {
        "machine learning": [
            {
                "question": "What is overfitting in machine learning, and name one way to reduce it.",
                "reference_answer": (
                    "Overfitting is when a model memorizes training noise and fails to generalize to unseen data. "
                    "It can be reduced using regularization, cross-validation, early stopping, or more data."
                ),
                "keywords": ["overfitting", "generalize", "regularization", "cross-validation"],
            },
            {
                "question": "Why do we split data into train, validation, and test sets?",
                "reference_answer": (
                    "The train set fits the model, validation tunes hyperparameters, and the test set estimates "
                    "final generalization on unseen data."
                ),
                "keywords": ["train", "validation", "test", "generalization"],
            },
        ],
        "gradient descent": [
            {
                "question": "Why does gradient descent use the negative gradient direction?",
                "reference_answer": (
                    "The gradient points to steepest increase in loss, so moving in the negative gradient direction "
                    "decreases loss most quickly for a small step."
                ),
                "keywords": ["negative gradient", "decrease loss", "steepest increase"],
            },
            {
                "question": "What happens if the learning rate is too high in gradient descent?",
                "reference_answer": (
                    "If learning rate is too high, updates overshoot minima and training can oscillate or diverge."
                ),
                "keywords": ["learning rate", "overshoot", "oscillate", "diverge"],
            },
        ],
        "backpropagation": [
            {
                "question": "What is backpropagation and why is the chain rule needed?",
                "reference_answer": (
                    "Backpropagation computes gradients of loss with respect to each weight. "
                    "The chain rule propagates derivatives through layers."
                ),
                "keywords": ["backpropagation", "gradients", "loss", "chain rule", "layers"],
            },
            {
                "question": "Why can vanishing gradients hurt deep neural network training?",
                "reference_answer": (
                    "When gradients become very small in early layers, those weights barely update, "
                    "so learning slows or stalls."
                ),
                "keywords": ["vanishing gradients", "small gradients", "early layers", "slow learning"],
            },
        ],
        "data analysis": [
            {
                "question": "What is the difference between descriptive and inferential statistics?",
                "reference_answer": (
                    "Descriptive statistics summarize observed data, while inferential statistics draw conclusions "
                    "about a population using samples."
                ),
                "keywords": ["descriptive", "inferential", "summarize", "population", "sample"],
            },
            {
                "question": "Why is data cleaning important before analysis?",
                "reference_answer": (
                    "Cleaning handles missing values, duplicates, and inconsistencies so analysis results are reliable."
                ),
                "keywords": ["data cleaning", "missing values", "duplicates", "reliable"],
            },
        ],
        "data engineering": [
            {
                "question": "What is ETL in data engineering?",
                "reference_answer": (
                    "ETL means Extract, Transform, and Load: collect data from sources, clean/reshape it, "
                    "then load it into a warehouse or lake for use."
                ),
                "keywords": ["extract", "transform", "load", "warehouse", "data lake"],
            },
            {
                "question": "Why is schema validation important in data pipelines?",
                "reference_answer": (
                    "Schema validation catches incompatible or malformed records early, preventing downstream failures "
                    "and poor data quality."
                ),
                "keywords": ["schema validation", "data quality", "pipeline", "downstream failures"],
            },
        ],
        "sql": [
            {
                "question": "What does a SQL JOIN do?",
                "reference_answer": (
                    "A JOIN combines rows from two tables using a related key so you can query connected data."
                ),
                "keywords": ["join", "two tables", "related key", "combine rows"],
            },
            {
                "question": "When would you use GROUP BY in SQL?",
                "reference_answer": (
                    "Use GROUP BY to aggregate rows by one or more columns, for example count or sum per category."
                ),
                "keywords": ["group by", "aggregate", "count", "sum", "category"],
            },
        ],
    }

    _CONCEPT_ALIAS_MAP: Dict[str, str] = {
        "ml": "machine learning",
        "ai": "machine learning",
        "data analyst": "data analysis",
        "data analytics": "data analysis",
        "analytics": "data analysis",
        "etl": "data engineering",
        "extract transform load": "data engineering",
        "extract, transform, load": "data engineering",
        "data pipelines": "data engineering",
        "pipeline": "data engineering",
        "pipelines": "data engineering",
        "sql query": "sql",
    }

    def __init__(self, db: Session):
        self.db = db

    @staticmethod
    def _normalize_concept(concept_name: str) -> str:
        value = (concept_name or "general").strip().lower()
        value = re.sub(r"[^a-z0-9\s\-/+]", " ", value)
        value = re.sub(r"\s+", " ", value).strip()
        if not value:
            return "general"

        if value in WeaknessAnalyzerService._CONCEPT_ALIAS_MAP:
            return WeaknessAnalyzerService._CONCEPT_ALIAS_MAP[value]

        noisy_markers = ["current level", "timeline", "beginner", "intermediate", "advanced"]
        if any(marker in value for marker in noisy_markers) and "machine learning" in value:
            return "machine learning"
        if any(marker in value for marker in noisy_markers) and "data analysis" in value:
            return "data analysis"
        if any(marker in value for marker in noisy_markers) and "data engineering" in value:
            return "data engineering"

        return value or "general"

    def _pick_quiz_concept(self, student_id: int, concept_name: Optional[str] = None) -> str:
        bank = self._QUIZ_QUESTION_BANK
        if concept_name:
            normalized = self._normalize_concept(concept_name)
            if normalized in bank:
                return normalized

        for weakness in self.get_weakest_concepts(student_id, limit=8):
            normalized = self._normalize_concept(weakness.concept_name)
            if normalized in bank:
                return normalized

        return random.choice(list(bank.keys()))

    def generate_quiz_question(self, student_id: int, concept_name: Optional[str] = None) -> Dict:
        student_exists = self.db.query(Student).filter(Student.id == student_id).first()
        if not student_exists:
            raise ValueError(f"Student {student_id} not found")

        concept_key = self._pick_quiz_concept(student_id=student_id, concept_name=concept_name)
        bank_items = self._QUIZ_QUESTION_BANK.get(concept_key, self._QUIZ_QUESTION_BANK["machine learning"])
        selected = random.choice(bank_items)
        return {
            "question_id": str(uuid.uuid4()),
            "concept_name": concept_key,
            "question": str(selected["question"]),
            "reference_answer": str(selected["reference_answer"]),
            "keywords": [str(item).strip().lower() for item in selected.get("keywords", [])],
        }

    @staticmethod
    def _normalize_answer_text(value: str) -> str:
        text = (value or "").strip().lower()
        text = re.sub(r"[^a-z0-9\s]", " ", text)
        return re.sub(r"\s+", " ", text).strip()

    @staticmethod
    def _answer_token_set(value: str) -> set:
        text = WeaknessAnalyzerService._normalize_answer_text(value)
        if not text:
            return set()
        stop_words = {
            "the",
            "and",
            "for",
            "with",
            "that",
            "this",
            "from",
            "into",
            "your",
            "you",
            "are",
            "was",
            "were",
            "have",
            "has",
            "had",
            "its",
            "can",
            "could",
            "would",
            "should",
            "about",
            "what",
            "when",
            "where",
            "which",
            "why",
            "how",
            "then",
            "than",
            "them",
            "they",
            "their",
            "there",
            "here",
            "also",
            "just",
            "very",
            "more",
            "most",
            "much",
            "many",
            "like",
            "only",
            "been",
            "being",
            "because",
            "through",
            "each",
        }
        tokens = [token for token in text.split() if len(token) > 2 and token not in stop_words]
        return set(tokens)

    @staticmethod
    def _keyword_match_ratio(student_answer: str, keywords: Optional[List[str]] = None) -> float:
        if not keywords:
            return 0.0
        student_text = WeaknessAnalyzerService._normalize_answer_text(student_answer)
        if not student_text:
            return 0.0

        normalized_keywords = [WeaknessAnalyzerService._normalize_answer_text(item) for item in keywords]
        valid_keywords = [item for item in normalized_keywords if item]
        if not valid_keywords:
            return 0.0

        matched = sum(1 for keyword in valid_keywords if keyword in student_text)
        return matched / len(valid_keywords)

    @staticmethod
    def _reference_overlap_ratio(student_answer: str, reference_answer: str) -> float:
        student_tokens = WeaknessAnalyzerService._answer_token_set(student_answer)
        reference_tokens = WeaknessAnalyzerService._answer_token_set(reference_answer)
        if not student_tokens or not reference_tokens:
            return 0.0
        return len(student_tokens & reference_tokens) / len(reference_tokens)

    def evaluate_quiz_answer(
        self,
        concept_name: str,
        student_answer: str,
        reference_answer: str,
        keywords: Optional[List[str]] = None,
    ) -> bool:
        student_text = self._normalize_answer_text(student_answer)
        reference_text = self._normalize_answer_text(reference_answer)
        if not student_text:
            return False
        if student_text == reference_text:
            return True

        keyword_ratio = self._keyword_match_ratio(student_answer, keywords)
        overlap_ratio = self._reference_overlap_ratio(student_answer, reference_answer)

        concept_key = self._normalize_concept(concept_name)
        strict_concepts = {"gradient descent", "backpropagation", "sql"}
        if concept_key in strict_concepts:
            return keyword_ratio >= 0.5 or (keyword_ratio >= 0.34 and overlap_ratio >= 0.45)

        return keyword_ratio >= 0.5 or overlap_ratio >= 0.55 or (keyword_ratio >= 0.34 and overlap_ratio >= 0.45)

    def analyze_generated_quiz_attempt(
        self,
        student_id: int,
        concept_name: str,
        question: str,
        student_answer: str,
        reference_answer: str,
        keywords: Optional[List[str]] = None,
    ) -> WeaknessAnalysisResult:
        del question  # question text is present for audit/input symmetry; scoring relies on concept + reference.
        is_correct = self.evaluate_quiz_answer(
            concept_name=concept_name,
            student_answer=student_answer,
            reference_answer=reference_answer,
            keywords=keywords,
        )
        return self.analyze_quiz_result(
            student_id=student_id,
            concept_name=concept_name,
            is_correct=is_correct,
            student_answer=student_answer,
            correct_answer=reference_answer,
        )

    def get_or_create_weakness(self, student_id: int, concept_name: str) -> WeaknessScore:
        concept_key = self._normalize_concept(concept_name)
        weakness = (
            self.db.query(WeaknessScore)
            .filter(WeaknessScore.student_id == student_id, WeaknessScore.concept_name == concept_key)
            .first()
        )

        if not weakness:
            weakness = WeaknessScore(
                student_id=student_id,
                concept_name=concept_key,
                weakness_score=0.0,
                times_seen=0,
                times_correct=0,
            )
            self.db.add(weakness)
            self.db.commit()
            self.db.refresh(weakness)

        return weakness

    def analyze_quiz_result(
        self,
        student_id: int,
        concept_name: str,
        is_correct: bool,
        student_answer: str = "",
        correct_answer: str = "",
    ) -> WeaknessAnalysisResult:
        student_exists = self.db.query(Student).filter(Student.id == student_id).first()
        if not student_exists:
            raise ValueError(f"Student {student_id} not found")

        weakness = self.get_or_create_weakness(student_id, concept_name)
        old_score = weakness.weakness_score

        weakness.update_from_quiz_result(is_correct)
        self.db.commit()

        misconception = None
        if not is_correct:
            misconception = self._detect_misconception(student_answer, correct_answer, concept_name)

        return WeaknessAnalysisResult(
            concept_name=weakness.concept_name,
            is_correct=is_correct,
            old_weakness_score=round(old_score, 3),
            new_weakness_score=round(weakness.weakness_score, 3),
            misconception_detected=misconception,
            learning_priority=self._calculate_learning_priority(weakness.weakness_score),
        )

    def get_weakest_concepts(self, student_id: int, limit: int = 5) -> List[WeaknessScore]:
        return (
            self.db.query(WeaknessScore)
            .filter(WeaknessScore.student_id == student_id)
            .order_by(desc(WeaknessScore.weakness_score), desc(WeaknessScore.times_seen))
            .limit(limit)
            .all()
        )

    def get_strength_areas(self, student_id: int, limit: int = 3) -> List[str]:
        rows = (
            self.db.query(WeaknessScore)
            .filter(WeaknessScore.student_id == student_id, WeaknessScore.times_seen >= 2)
            .order_by(WeaknessScore.weakness_score.asc(), WeaknessScore.times_correct.desc())
            .limit(limit)
            .all()
        )
        return [item.concept_name for item in rows if item.weakness_score <= 0.35]

    def _detect_misconception(self, student_answer: str, correct_answer: str, concept_name: str) -> Optional[str]:
        student = (student_answer or "").strip().lower()
        correct = (correct_answer or "").strip().lower()
        concept = self._normalize_concept(concept_name)

        if not student:
            return f"No reasoning provided for {concept}."
        if student == correct:
            return None
        if any(token in student for token in ["always", "never"]):
            return f"Over-generalization in {concept}."
        if any(token in student for token in ["random", "guess", "maybe"]):
            return f"Uncertain causal reasoning in {concept}."

        return f"Core concept mismatch in {concept}."

    @staticmethod
    def _calculate_learning_priority(weakness_score: float) -> str:
        if weakness_score >= 0.75:
            return "critical"
        if weakness_score >= 0.5:
            return "high"
        if weakness_score >= 0.25:
            return "medium"
        return "low"

    def explain_mistake(
        self,
        student_id: int,
        concept: str,
        student_answer: str,
        correct_answer: str,
        question: Optional[str] = None,
    ) -> MistakeExplanation:
        student_exists = self.db.query(Student).filter(Student.id == student_id).first()
        if not student_exists:
            raise ValueError(f"Student {student_id} not found")

        concept_key = self._normalize_concept(concept)
        misconception = self._detect_misconception(student_answer, correct_answer, concept_key) or f"Gap in {concept_key}."

        why_wrong = (
            f"Your answer focuses on '{student_answer}', but it misses the key idea needed in {concept_key}. "
            "The response does not align with the underlying principle used to solve this type of problem."
        )
        correct_explanation = (
            f"A stronger answer is '{correct_answer}'. The main idea is to reason from first principles in {concept_key} "
            "before applying formulas or shortcuts."
        )
        guiding_question = self._build_guiding_question(concept_key, question)

        return MistakeExplanation(
            student_id=student_id,
            concept=concept_key,
            misconception_identified=misconception,
            why_wrong=why_wrong,
            correct_explanation=correct_explanation,
            learning_tips=[
                "Restate the concept in your own words before solving.",
                "Identify what assumption your answer depends on.",
                "Test your reasoning on one simpler example first.",
            ],
            related_concept="prerequisite foundations",
            guiding_question=guiding_question,
        )

    @staticmethod
    def _build_guiding_question(concept: str, question: Optional[str]) -> str:
        if question:
            return f"In '{question}', which step in {concept} determines whether your approach is valid?"
        return f"What evidence would convince you that your current reasoning in {concept} is correct?"


class ResumeMentorService:
    """Resume upload analyzer with Socratic mentoring guidance."""

    _SECTION_ALIASES: Dict[str, List[str]] = {
        "summary": ["summary", "profile", "objective", "about me"],
        "experience": ["experience", "work experience", "employment", "professional experience", "internship"],
        "projects": ["projects", "project experience", "academic projects"],
        "skills": ["skills", "technical skills", "core skills", "technologies", "tooling"],
        "education": ["education", "academic background", "qualifications"],
        "certifications": ["certifications", "certificates", "licenses"],
    }

    _REQUIRED_SECTIONS: List[str] = ["summary", "experience", "projects", "skills", "education"]
    _ACTION_VERBS = {
        "built",
        "developed",
        "improved",
        "optimized",
        "designed",
        "implemented",
        "delivered",
        "led",
        "created",
        "reduced",
        "increased",
        "automated",
        "deployed",
        "analyzed",
        "achieved",
    }

    def __init__(self, db: Session):
        self.db = db

    @staticmethod
    def _safe_decode(raw: bytes) -> str:
        for encoding in ("utf-8", "utf-16", "latin-1"):
            try:
                return raw.decode(encoding)
            except Exception:
                continue
        return raw.decode("utf-8", errors="ignore")

    def _extract_text(self, file_name: str, raw_bytes: bytes) -> str:
        lower_name = (file_name or "").lower()
        if lower_name.endswith(".txt") or lower_name.endswith(".md"):
            return self._safe_decode(raw_bytes)

        if lower_name.endswith(".pdf"):
            try:
                from pypdf import PdfReader  # type: ignore
            except Exception as e:
                raise ValueError("PDF parsing requires `pypdf`. Add it to requirements.") from e

            from io import BytesIO

            reader = PdfReader(BytesIO(raw_bytes))
            pages = [page.extract_text() or "" for page in reader.pages]
            return "\n".join(pages)

        if lower_name.endswith(".docx"):
            try:
                from docx import Document  # type: ignore
            except Exception as e:
                raise ValueError("DOCX parsing requires `python-docx`. Add it to requirements.") from e

            from io import BytesIO

            document = Document(BytesIO(raw_bytes))
            return "\n".join([paragraph.text for paragraph in document.paragraphs])

        raise ValueError("Unsupported file type. Upload .pdf, .docx, .txt, or .md.")

    @classmethod
    def _normalize_line(cls, line: str) -> str:
        normalized = re.sub(r"[^a-z0-9\s]", " ", (line or "").strip().lower())
        return re.sub(r"\s+", " ", normalized).strip()

    @classmethod
    def _match_section_name(cls, line: str) -> Optional[str]:
        normalized = cls._normalize_line(line)
        if not normalized:
            return None
        for canonical, aliases in cls._SECTION_ALIASES.items():
            for alias in aliases:
                if normalized == alias or normalized.startswith(f"{alias} "):
                    return canonical
        return None

    @classmethod
    def _split_sections(cls, text: str) -> Dict[str, str]:
        lines = [line.strip() for line in (text or "").splitlines()]
        sections: Dict[str, List[str]] = {}
        active_section = "general"
        sections.setdefault(active_section, [])

        for line in lines:
            if not line:
                continue
            section_name = cls._match_section_name(line.rstrip(":"))
            if section_name:
                active_section = section_name
                sections.setdefault(active_section, [])
                continue
            sections.setdefault(active_section, []).append(line)

        return {name: "\n".join(content).strip() for name, content in sections.items() if "\n".join(content).strip()}

    @staticmethod
    def _is_bullet(line: str) -> bool:
        stripped = (line or "").strip()
        return bool(re.match(r"^(\-|\*|•|\d+\.)\s+", stripped))

    @staticmethod
    def _has_number(text: str) -> bool:
        return bool(re.search(r"\d", text or ""))

    def _analyze_section(self, section_name: str, content: str) -> Tuple[ResumeSectionAnalysis, List[ResumeIssue]]:
        text = (content or "").strip()
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        score = 1.0
        findings: List[str] = []
        mentoring_questions: List[str] = []
        issues: List[ResumeIssue] = []

        if len(text) < 80 and section_name in self._REQUIRED_SECTIONS:
            score -= 0.3
            finding = "Section is too brief and lacks enough detail."
            findings.append(finding)
            mentoring_questions.append(f"What concrete details can you add to strengthen your {section_name} section?")
            issues.append(
                ResumeIssue(
                    issue_type="poor structure",
                    severity="medium",
                    section_name=section_name,
                    evidence=finding,
                    mentoring_question=mentoring_questions[-1],
                )
            )

        if section_name in {"experience", "projects"}:
            bullets = [line for line in lines if self._is_bullet(line)]
            if not bullets:
                score -= 0.3
                evidence = "No bullet points found for achievements."
                question = "Can you rewrite your work into bullet points with action + impact?"
                findings.append(evidence)
                mentoring_questions.append(question)
                issues.append(
                    ResumeIssue(
                        issue_type="weak bullet points",
                        severity="high",
                        section_name=section_name,
                        evidence=evidence,
                        mentoring_question=question,
                    )
                )
            else:
                quantified = [item for item in bullets if self._has_number(item)]
                if not quantified:
                    score -= 0.3
                    evidence = "Bullets describe tasks, but none quantify impact."
                    question = "Can you quantify impact for each bullet with metrics like %, time saved, or users affected?"
                    findings.append(evidence)
                    mentoring_questions.append(question)
                    issues.append(
                        ResumeIssue(
                            issue_type="unclear achievements",
                            severity="high",
                            section_name=section_name,
                            evidence=evidence,
                            mentoring_question=question,
                        )
                    )

                weak_bullets = []
                for item in bullets:
                    normalized = self._normalize_line(item)
                    if not any(f" {verb} " in f" {normalized} " for verb in self._ACTION_VERBS):
                        weak_bullets.append(item)
                if weak_bullets:
                    score -= 0.15
                    evidence = "Some bullets start without strong action verbs."
                    question = "How can you start each bullet with a strong action verb like built, improved, or optimized?"
                    findings.append(evidence)
                    mentoring_questions.append(question)
                    issues.append(
                        ResumeIssue(
                            issue_type="weak bullet points",
                            severity="medium",
                            section_name=section_name,
                            evidence=evidence,
                            mentoring_question=question,
                        )
                    )

        if section_name == "skills":
            tokens = re.split(r"[,/\n|]", text)
            unique_skills = [token.strip() for token in tokens if token.strip()]
            if len(unique_skills) < 6:
                score -= 0.35
                evidence = "Skills section appears too short for a technical resume."
                question = "Which core tools, languages, and frameworks are missing that match your target role?"
                findings.append(evidence)
                mentoring_questions.append(question)
                issues.append(
                    ResumeIssue(
                        issue_type="missing skills",
                        severity="high",
                        section_name=section_name,
                        evidence=evidence,
                        mentoring_question=question,
                    )
                )

        if section_name == "summary":
            lower = text.lower()
            generic_markers = ["hardworking", "passionate", "seeking opportunity", "quick learner"]
            if any(marker in lower for marker in generic_markers):
                score -= 0.2
                evidence = "Summary uses generic phrases without role-specific evidence."
                question = "Can you make the summary role-specific with one measurable achievement?"
                findings.append(evidence)
                mentoring_questions.append(question)
                issues.append(
                    ResumeIssue(
                        issue_type="poor structure",
                        severity="medium",
                        section_name=section_name,
                        evidence=evidence,
                        mentoring_question=question,
                    )
                )

        if section_name == "education":
            if not re.search(r"(19|20)\d{2}", text):
                score -= 0.15
                evidence = "Education section has no visible graduation year."
                question = "Can you add graduation year and degree details for clarity?"
                findings.append(evidence)
                mentoring_questions.append(question)
                issues.append(
                    ResumeIssue(
                        issue_type="poor structure",
                        severity="low",
                        section_name=section_name,
                        evidence=evidence,
                        mentoring_question=question,
                    )
                )

        score = max(0.0, min(1.0, round(score, 2)))
        if not findings:
            findings.append("Section is reasonably clear.")
            mentoring_questions.append(f"What one improvement would make this {section_name} section even stronger?")

        return (
            ResumeSectionAnalysis(
                section_name=section_name,
                score=score,
                findings=findings,
                mentoring_questions=mentoring_questions,
            ),
            issues,
        )

    @classmethod
    def _missing_sections(cls, sections: Dict[str, str]) -> List[str]:
        return [name for name in cls._REQUIRED_SECTIONS if not sections.get(name)]

    @staticmethod
    def _overall_assessment(avg_score: float, issue_count: int, missing_count: int) -> str:
        if missing_count >= 2 or avg_score < 0.45:
            return "Resume needs major improvement before applying."
        if issue_count >= 4 or avg_score < 0.7:
            return "Resume is promising but needs targeted improvements."
        return "Resume is solid; refine impact statements to stand out."

    def analyze_resume(self, file_name: str, raw_bytes: bytes) -> ResumeMentorResponse:
        if not raw_bytes:
            raise ValueError("Uploaded file is empty.")
        if len(raw_bytes) > 6 * 1024 * 1024:
            raise ValueError("File too large. Please upload a file up to 6 MB.")

        resume_text = self._extract_text(file_name=file_name, raw_bytes=raw_bytes)
        if len((resume_text or "").strip()) < 80:
            raise ValueError("Could not extract enough resume content. Please upload a clearer file.")

        sections = self._split_sections(resume_text)
        missing_sections = self._missing_sections(sections)

        section_analysis: List[ResumeSectionAnalysis] = []
        issues: List[ResumeIssue] = []
        strengths: List[str] = []

        for section_name, content in sections.items():
            analysis, section_issues = self._analyze_section(section_name, content)
            section_analysis.append(analysis)
            issues.extend(section_issues)

            if analysis.score >= 0.8 and section_name in self._REQUIRED_SECTIONS:
                strengths.append(f"{section_name.title()} section is clear and reasonably detailed.")

            if section_name in {"experience", "projects"}:
                bullet_lines = [line for line in content.splitlines() if self._is_bullet(line)]
                quantified_lines = [line for line in bullet_lines if self._has_number(line)]
                if quantified_lines:
                    strengths.append(f"{section_name.title()} includes measurable impact in some bullets.")

        for section_name in missing_sections:
            issues.append(
                ResumeIssue(
                    issue_type="poor structure",
                    severity="high",
                    section_name=section_name,
                    evidence=f"Missing expected section: {section_name}.",
                    mentoring_question=f"How will you add a concise {section_name} section to improve resume structure?",
                )
            )

        if missing_sections:
            strengths = strengths or ["Resume content exists, but structure is incomplete."]

        weaknesses = [f"{issue.issue_type} in {issue.section_name}: {issue.evidence}" for issue in issues[:10]]

        unique_questions = []
        seen = set()
        for issue in issues:
            question = issue.mentoring_question.strip()
            if question and question not in seen:
                unique_questions.append(question)
                seen.add(question)
            if len(unique_questions) >= 8:
                break

        if not unique_questions:
            unique_questions = [
                "What role are you targeting, and which achievements best prove fit for that role?",
                "Can you quantify one project outcome with a measurable result?",
            ]

        average_score = round(
            sum(item.score for item in section_analysis) / max(1, len(section_analysis)),
            2,
        )
        overall = self._overall_assessment(
            avg_score=average_score,
            issue_count=len(issues),
            missing_count=len(missing_sections),
        )

        detected_keywords, missing_keywords = keyword_gap_analysis(
            resume_text=resume_text,
            important_keywords=AI_DATA_KEYWORDS,
        )
        score_breakdown = calculate_resume_score(
            resume_text=resume_text,
            sections=sections,
            detected_keywords=detected_keywords,
            important_keywords=AI_DATA_KEYWORDS,
        )
        suggestions = improvement_suggestions(
            resume_text=resume_text,
            sections=sections,
            missing_sections=missing_sections,
            missing_keywords=missing_keywords,
            score=score_breakdown,
        )

        return ResumeMentorResponse(
            file_name=file_name,
            overall_assessment=overall,
            resume_score=score_breakdown.total,
            detected_keywords=detected_keywords,
            missing_keywords=missing_keywords,
            detected_sections=sorted(list(sections.keys())),
            missing_sections=missing_sections,
            improvement_suggestions=suggestions,
            strengths=strengths[:8],
            weaknesses=weaknesses[:10],
            issues=issues[:12],
            section_analysis=section_analysis,
            mentoring_advice=unique_questions,
        )


class MentorAIService:
    """Generate adaptive Socratic mentor responses."""
    _llm_backoff_until: float = 0.0

    def __init__(self, db: Session):
        self.db = db
        self.profile_service = StudentProfileService(db)
        self.weakness_service = WeaknessAnalyzerService(db)

    def generate_response(
        self,
        student_id: int,
        query: str,
        focus_concept: Optional[str] = None,
        context: Optional[Dict] = None,
    ) -> Dict:
        profile = self.profile_service.get_profile(student_id)
        if not profile:
            raise ValueError(f"Profile not found for student {student_id}")

        latest_response = self._latest_response(student_id)
        is_follow_up_turn = self._is_follow_up_turn(query=query, latest_response=latest_response)

        target_concept = self._resolve_target_concept(
            student_id=student_id,
            query=query,
            focus_concept=focus_concept,
        )
        target_concept = self._carry_forward_concept_if_needed(
            query=query,
            resolved_concept=target_concept,
            latest_response=latest_response,
            is_follow_up_turn=is_follow_up_turn,
        )
        weakness = self.weakness_service.get_or_create_weakness(student_id, target_concept)

        feedback_bias = self._recent_feedback_bias(student_id)
        explanation_style = self._determine_explanation_style(
            profile.confidence_level,
            weakness.weakness_score,
            profile.preferred_difficulty,
            feedback_bias,
        )

        response_text = self._generate_socratic_response(
            student_id=student_id,
            query=query,
            concept=target_concept,
            style=explanation_style,
            context=context or profile.learning_style_summary,
            latest_response=latest_response,
            is_follow_up_turn=is_follow_up_turn,
        )
        response_text = self._avoid_repetitive_reply(
            query=query,
            concept=target_concept,
            response_text=response_text,
            latest_response=latest_response,
            is_follow_up_turn=is_follow_up_turn,
        )
        follow_up = self._generate_guiding_question(target_concept, explanation_style)

        response_id = self._store_response(
            student_id=student_id,
            query=query,
            response=response_text,
            style=explanation_style,
            concept=target_concept,
            confidence=profile.confidence_level,
        )

        return {
            "response_id": response_id,
            "response": response_text,
            "explanation_style": explanation_style,
            "target_concept": target_concept,
            "follow_up_question": follow_up,
        }

    def _resolve_target_concept(
        self,
        student_id: int,
        query: str,
        focus_concept: Optional[str] = None,
    ) -> str:
        if focus_concept:
            return self._normalize_topic_alias(focus_concept)

        inferred = self._infer_concept(query)
        extracted_topic = self._extract_topic_from_query(query)
        query_l = (query or "").lower()
        if self._is_skills_query(query_l):
            if any(token in query_l for token in ["machine learning", " ml ", " ai", "ai ", "artificial intelligence"]):
                return "machine learning"
            if any(token in query_l for token in ["data analyst", "data analysis", "analytics", "business intelligence"]):
                return "data analysis"
            if "backend" in query_l:
                return "backend engineering"
        continuation_tokens = [
            "current level",
            "timeline",
            "month",
            "months",
            "week",
            "weeks",
            "beginner",
            "intermediate",
            "advanced",
        ]
        is_continuation = any(token in query_l for token in continuation_tokens)
        continuation_topic_markers = {
            "beginner",
            "intermediate",
            "advanced",
            "your beginner",
            "your intermediate",
            "your advanced",
            "current level",
            "timeline",
        }
        extracted_is_meta = (
            extracted_topic in continuation_topic_markers
            or extracted_topic.startswith("your ")
        )
        has_explicit_topic = (
            inferred != "general"
            or (extracted_topic and extracted_topic != "general" and not extracted_is_meta)
        )

        if is_continuation and not has_explicit_topic:
            recent = (
                self.db.query(MentorResponse)
                .filter(MentorResponse.student_id == student_id)
                .order_by(desc(MentorResponse.created_at))
                .limit(3)
                .all()
            )

            for row in recent:
                previous_query = (row.query or "").lower()
                asked_roadmap = any(token in previous_query for token in ["roadmap", "plan", "path", "become", "career", "learn"])
                if asked_roadmap and row.target_concept and row.target_concept != "general":
                    return row.target_concept

        if extracted_topic and extracted_topic != "general":
            return self._normalize_topic_alias(extracted_topic)

        if inferred != "general":
            return inferred

        if extracted_topic and extracted_topic != "general":
            return self._normalize_topic_alias(extracted_topic)
        if not is_continuation:
            return inferred

        recent = (
            self.db.query(MentorResponse)
            .filter(MentorResponse.student_id == student_id)
            .order_by(desc(MentorResponse.created_at))
            .limit(3)
            .all()
        )

        for row in recent:
            previous_query = (row.query or "").lower()
            asked_roadmap = any(token in previous_query for token in ["roadmap", "plan", "path", "become", "career"])
            if asked_roadmap and row.target_concept and row.target_concept != "general":
                return row.target_concept

        return inferred

    def _recent_feedback_bias(self, student_id: int) -> str:
        rows = (
            self.db.query(Feedback.feedback_type)
            .filter(Feedback.student_id == student_id)
            .order_by(desc(Feedback.created_at))
            .limit(5)
            .all()
        )
        if not rows:
            return "neutral"

        hard_votes = sum(1 for row in rows if row.feedback_type in (FeedbackType.TOO_HARD, FeedbackType.UNCLEAR))
        easy_votes = sum(1 for row in rows if row.feedback_type == FeedbackType.TOO_EASY)

        if hard_votes > easy_votes:
            return "simplify"
        if easy_votes > hard_votes:
            return "deepen"
        return "neutral"

    @staticmethod
    def _determine_explanation_style(
        confidence: float,
        weakness: float,
        preferred_difficulty: DifficultyLevel,
        feedback_bias: str,
    ) -> str:
        if feedback_bias == "simplify":
            return "simple"
        if feedback_bias == "deepen":
            return "deep"

        if weakness >= 0.6 or confidence <= 0.35:
            return "simple"
        if weakness >= 0.3 or preferred_difficulty == DifficultyLevel.MEDIUM:
            return "conceptual"
        return "deep"

    @staticmethod
    def _infer_concept(query: str) -> str:
        query_l = (query or "").lower()
        keyword_map = {
            "gradient descent": ["gradient", "descent", "learning rate", "loss"],
            "backpropagation": ["backprop", "backward", "weights", "neural"],
            "machine learning": [
                "machine learning",
                "ml",
                "ai",
                "artificial intelligence",
                "supervised",
                "unsupervised",
                "model training",
                "feature engineering",
                "overfitting",
                "underfitting",
                "regularization",
                "classification",
                "regression",
            ],
            "statistics": ["probability", "variance", "mean", "hypothesis"],
            "linear algebra": ["matrix", "vector", "eigen", "determinant"],
            "python": ["python", "list", "dict", "function", "class"],
            "sql": ["sql", "join", "select", "where", "group by"],
            "data analysis": ["data analyst", "data analysis", "analysis", "dashboard", "kpi", "analytics", "bi"],
            "data engineering": [
                "data engineering",
                "etl",
                "extract transform load",
                "extract, transform, load",
                "data pipeline",
                "data pipelines",
                "data warehouse",
                "data lake",
                "orchestration",
                "airflow",
                "spark",
            ],
            "data structures": [
                "data structure",
                "array",
                "linked list",
                "stack",
                "queue",
                "tree",
                "graph",
                "hashmap",
                "hashing",
                "hash table",
            ],
            "algorithms": ["algorithm", "complexity", "big o", "recursion", "dynamic programming", "greedy", "sorting"],
            "system design": ["system design", "scalability", "load balancer", "caching", "distributed"],
            "api design": ["api", "rest", "endpoint", "request", "response", "http"],
        }

        best = "general"
        best_hits = 0
        for concept, words in keyword_map.items():
            hits = sum(1 for word in words if word in query_l)
            if hits > best_hits:
                best = concept
                best_hits = hits
        return best

    def _latest_response(self, student_id: int) -> Optional[MentorResponse]:
        return (
            self.db.query(MentorResponse)
            .filter(MentorResponse.student_id == student_id)
            .order_by(desc(MentorResponse.created_at))
            .first()
        )

    @staticmethod
    def _content_tokens(text: str) -> set:
        raw_tokens = re.findall(r"[a-z0-9]+", (text or "").lower())
        stop_words = {
            "the",
            "and",
            "for",
            "with",
            "that",
            "this",
            "from",
            "into",
            "your",
            "you",
            "are",
            "was",
            "were",
            "have",
            "has",
            "had",
            "its",
            "can",
            "could",
            "would",
            "should",
            "about",
            "what",
            "when",
            "where",
            "which",
            "why",
            "how",
            "then",
            "than",
            "them",
            "they",
            "their",
            "there",
            "here",
            "also",
            "just",
            "very",
            "more",
            "most",
            "much",
            "many",
            "like",
            "only",
            "been",
            "being",
            "because",
            "through",
            "each",
            "step",
        }
        return {token for token in raw_tokens if len(token) > 2 and token not in stop_words}

    @staticmethod
    def _token_overlap_ratio(text_a: str, text_b: str) -> float:
        tokens_a = MentorAIService._content_tokens(text_a)
        tokens_b = MentorAIService._content_tokens(text_b)
        if not tokens_a or not tokens_b:
            return 0.0
        return len(tokens_a & tokens_b) / max(1, len(tokens_a))

    @staticmethod
    def _jaccard_similarity(text_a: str, text_b: str) -> float:
        tokens_a = MentorAIService._content_tokens(text_a)
        tokens_b = MentorAIService._content_tokens(text_b)
        union = tokens_a | tokens_b
        if not union:
            return 0.0
        return len(tokens_a & tokens_b) / len(union)

    def _is_follow_up_turn(self, query: str, latest_response: Optional[MentorResponse]) -> bool:
        if not latest_response:
            return False

        query_l = (query or "").strip().lower()
        if not query_l or "?" in query_l:
            return False

        token_count = len(query_l.split())
        if token_count == 0 or token_count > 30:
            return False

        previous_concept = (latest_response.target_concept or "").lower()
        mentions_previous_concept = bool(
            previous_concept and previous_concept != "general" and previous_concept in query_l
        )
        reflective_phrases = [
            "i think",
            "i guess",
            "it means",
            "so it",
            "from my understanding",
            "my understanding",
            "basically",
            "in short",
            "it is",
            "this is",
        ]
        sounds_like_reflection = any(phrase in query_l for phrase in reflective_phrases)
        overlap_with_previous = self._token_overlap_ratio(query_l, latest_response.response or "") >= 0.35
        overlap_with_previous_query = self._token_overlap_ratio(query_l, latest_response.query or "") >= 0.4

        return (
            mentions_previous_concept
            or sounds_like_reflection
            or overlap_with_previous
            or overlap_with_previous_query
        )

    def _carry_forward_concept_if_needed(
        self,
        query: str,
        resolved_concept: str,
        latest_response: Optional[MentorResponse],
        is_follow_up_turn: bool,
    ) -> str:
        if resolved_concept and resolved_concept != "general":
            return resolved_concept
        if not latest_response:
            return resolved_concept

        previous_concept = (latest_response.target_concept or "").strip()
        if not previous_concept or previous_concept == "general":
            return resolved_concept

        if is_follow_up_turn:
            return previous_concept

        overlap_with_previous_query = self._token_overlap_ratio(query, latest_response.query or "")
        if overlap_with_previous_query >= 0.5:
            return previous_concept
        return resolved_concept

    def _avoid_repetitive_reply(
        self,
        query: str,
        concept: str,
        response_text: str,
        latest_response: Optional[MentorResponse],
        is_follow_up_turn: bool,
    ) -> str:
        if not response_text or not latest_response or not latest_response.response:
            return response_text

        similarity = self._jaccard_similarity(response_text, latest_response.response)
        threshold = 0.55 if is_follow_up_turn else 0.82
        if similarity < threshold:
            return response_text

        concept_label = concept if concept and concept != "general" else (latest_response.target_concept or "this topic")
        if concept_label == "data engineering":
            return (
                "Good follow-up. You already captured the ETL idea.\n"
                "Now push one level deeper:\n"
                "1. Why transform before load (quality checks, schema mapping, deduplication).\n"
                "2. Where the data lands (warehouse/lake) and how success is validated.\n"
                "3. One real tool per stage (for example, Airflow + Spark + BigQuery/Snowflake).\n\n"
                "Your turn: for event logs with missing timestamps, what transformation would you apply before loading?"
            )

        return (
            f"Good follow-up. You already captured the core of {concept_label}.\n"
            "Instead of repeating the same definition, add:\n"
            "1. Why each step matters.\n"
            "2. One concrete example or tool.\n"
            "3. One common pitfall and how to detect it.\n\n"
            f"Can you apply that to one small {concept_label} example?"
        )

    @staticmethod
    def _clean_topic_phrase(value: str) -> str:
        phrase = (value or "").strip().lower()
        phrase = re.sub(r"[^a-z0-9\s\-/+]", " ", phrase)
        phrase = re.sub(r"\s+", " ", phrase).strip()
        if not phrase:
            return "general"

        filler = {
            "please",
            "can",
            "you",
            "me",
            "give",
            "the",
            "a",
            "an",
            "to",
            "for",
            "of",
            "about",
            "in",
            "on",
            "is",
            "are",
            "do",
            "does",
            "how",
            "why",
            "what",
            "detailed",
            "important",
            "happen",
            "happens",
            "current",
            "level",
            "timeline",
            "learn",
            "study",
            "from",
            "scratch",
            "we",
            "use",
            "become",
            "month",
            "months",
            "required",
            "skill",
            "skills",
            "prerequisite",
            "prerequisites",
        }
        words = [word for word in phrase.split() if word not in filler and not word.isdigit()]
        if not words:
            return "general"
        return " ".join(words[:5])

    @staticmethod
    def _normalize_topic_alias(topic: str) -> str:
        cleaned = MentorAIService._clean_topic_phrase(topic)
        alias_map = {
            "ml": "machine learning",
            "ai": "machine learning",
            "data analyst": "data analysis",
            "data analytics": "data analysis",
            "analytics": "data analysis",
            "sql query": "sql",
        }
        return alias_map.get(cleaned, cleaned)

    @staticmethod
    def _extract_topic_from_query(query: str) -> str:
        text = (query or "").strip().lower()
        if not text:
            return "general"

        cleaned = re.sub(r"[?!.,]", " ", text)
        cleaned = re.sub(r"\s+", " ", cleaned).strip()

        patterns = [
            r"^(?:what is|what are|define|explain)\s+(.+)$",
            r"^(?:why does|why do|why is|why are|why)\s+(.+)$",
            r"^(?:how does|how do|how can|how to|how)\s+(.+)$",
            r"(?:roadmap|road map|plan|path)\s+(?:for|to become)\s+(.+)$",
            r"(?:learn|study)\s+(.+)$",
            r"(?:difference between|compare)\s+(.+?)\s+(?:and|vs)\s+(.+)$",
        ]

        for pattern in patterns:
            match = re.search(pattern, cleaned)
            if not match:
                continue

            if len(match.groups()) == 1:
                return MentorAIService._clean_topic_phrase(match.group(1))
            left = MentorAIService._clean_topic_phrase(match.group(1))
            right = MentorAIService._clean_topic_phrase(match.group(2))
            if left != "general" and right != "general":
                return f"{left} vs {right}"

        if "for " in cleaned:
            tail = cleaned.split("for ", 1)[1]
            return MentorAIService._clean_topic_phrase(tail)

        return "general"

    def _generate_socratic_response(
        self,
        student_id: int,
        query: str,
        concept: str,
        style: str,
        context: Dict,
        latest_response: Optional[MentorResponse] = None,
        is_follow_up_turn: bool = False,
    ) -> str:
        if self._is_skills_query((query or "").lower()):
            return self._skills_requirement_response(query=query, concept=concept, context=context)

        llm_text = self._try_llm_response(
            student_id=student_id,
            query=query,
            concept=concept,
            style=style,
            context=context,
            latest_response=latest_response,
            is_follow_up_turn=is_follow_up_turn,
        )
        if llm_text:
            return llm_text
        return self._local_socratic_response(
            query=query,
            concept=concept,
            style=style,
            latest_response=latest_response,
            is_follow_up_turn=is_follow_up_turn,
        )

    @staticmethod
    def _is_skills_query(query_l: str) -> bool:
        skill_tokens = [
            "what skills",
            "which skills",
            "required skills",
            "skills required",
            "skill set",
            "prerequisite",
            "prerequisites",
            "skills for",
            "need to learn",
        ]
        return any(token in query_l for token in skill_tokens)

    def _skills_requirement_response(self, query: str, concept: str, context: Dict) -> str:
        query_l = (query or "").lower()
        normalized_concept = self._normalize_topic_alias(concept)
        known_skills: List[str] = []
        if isinstance(context, dict):
            raw = context.get("skills")
            if isinstance(raw, list):
                known_skills = [str(item).strip() for item in raw if str(item).strip()]

        if any(token in query_l for token in ["ai", "artificial intelligence", "machine learning", "ml"]):
            normalized_concept = "machine learning"
        elif any(token in query_l for token in ["data analyst", "data analysis", "analytics", "business intelligence"]):
            normalized_concept = "data analysis"

        if normalized_concept == "machine learning":
            intro = "Required skills for AI/ML:"
            core = (
                "1. Programming: Python, data structures, Git, clean coding\n"
                "2. Math: linear algebra, probability, statistics, basic calculus\n"
                "3. ML fundamentals: supervised/unsupervised learning, model evaluation, bias-variance\n"
                "4. Data skills: pandas, SQL, feature engineering, data cleaning\n"
                "5. Model tooling: scikit-learn plus TensorFlow/PyTorch basics\n"
                "6. Portfolio: 2-3 end-to-end projects with clear metrics and business impact"
            )
        elif normalized_concept == "data analysis":
            intro = "Required skills for Data Analyst:"
            core = (
                "1. SQL: joins, group by, window functions\n"
                "2. Excel/Sheets: pivots, lookups, data cleaning\n"
                "3. Statistics: distributions, hypothesis testing, confidence intervals\n"
                "4. BI: Power BI or Tableau dashboards and storytelling\n"
                "5. Python (highly useful): pandas + visualization\n"
                "6. Communication: convert findings into clear recommendations"
            )
        else:
            intro = f"Required skills for {normalized_concept}:"
            core = (
                "1. Fundamentals and core terminology\n"
                "2. Practical tools used in real projects\n"
                "3. Problem-solving patterns and trade-offs\n"
                "4. Portfolio projects with measurable outcomes\n"
                "5. Interview-style practice and clear communication"
            )

        known_line = ""
        if known_skills:
            top_known = ", ".join(known_skills[:4])
            known_line = (
                f"\n\nYou already have: {top_known}.\n"
                "Focus next on the missing math/statistics and project depth."
            )

        return (
            f"{intro}\n{core}"
            f"{known_line}\n\n"
            "Tell me your level and timeline, and I will convert this into a week-by-week plan."
        )

    def _select_model(self) -> str:
        configured = os.getenv("OPENAI_API_MODEL")
        if configured:
            return configured

        base_url = (os.getenv("OPENAI_API_BASE") or "").lower()
        if "openrouter" in base_url:
            return "openrouter/auto"
        return "gpt-4o-mini"

    def _model_candidates(self) -> List[str]:
        primary = self._select_model()
        base_url = (os.getenv("OPENAI_API_BASE") or "").lower()

        candidates = [primary]
        if "openrouter" in base_url:
            fallback_models = [
                "openrouter/auto",
                "openai/gpt-4o-mini",
                "anthropic/claude-3.5-haiku",
                "meta-llama/llama-3.1-8b-instruct:free",
            ]
            for model in fallback_models:
                if model not in candidates:
                    candidates.append(model)
        return candidates

    def _recent_context(self, student_id: int, limit: int = 3) -> str:
        rows = (
            self.db.query(MentorResponse.query, MentorResponse.response, MentorResponse.target_concept)
            .filter(MentorResponse.student_id == student_id)
            .order_by(desc(MentorResponse.created_at))
            .limit(limit)
            .all()
        )
        if not rows:
            return "No prior mentor interactions."

        lines = []
        for idx, row in enumerate(reversed(rows), start=1):
            q_text = row[0] if isinstance(row, tuple) or isinstance(row, list) else row.query
            r_text = row[1] if isinstance(row, tuple) or isinstance(row, list) else row.response
            c_text = row[2] if isinstance(row, tuple) or isinstance(row, list) else row.target_concept
            lines.append(f"{idx}) Query: {q_text}")
            lines.append(f"   Concept: {c_text}")
            lines.append(f"   Mentor summary: {(r_text or '')[:220]}")
        return "\n".join(lines)

    def _try_llm_response(
        self,
        student_id: int,
        query: str,
        concept: str,
        style: str,
        context: Dict,
        latest_response: Optional[MentorResponse] = None,
        is_follow_up_turn: bool = False,
    ) -> Optional[str]:
        api_key = os.getenv("OPENAI_API_KEY")
        api_base = os.getenv("OPENAI_API_BASE")
        if not api_key:
            return None
        if time.time() < MentorAIService._llm_backoff_until:
            return None

        candidate_models = self._model_candidates()
        style_instruction = {
            "simple": "Use beginner-friendly language, short steps, and one small concrete example.",
            "conceptual": "Give a structured conceptual explanation with intuition + practical example.",
            "deep": "Give rigorous explanation with equations/assumptions when useful.",
        }.get(style, "Give a clear structured explanation.")

        system_prompt = (
            "You are an expert AI mentor.\n"
            "Always answer the student's actual question directly first.\n"
            "Then add 1-2 Socratic follow-up questions.\n"
            "Do not give generic placeholders. Keep output practical and relevant.\n"
            "If the student follow-up overlaps with your previous answer, avoid repeating the same explanation."
        )
        previous_query = (latest_response.query if latest_response else "None")
        previous_answer_summary = ((latest_response.response or "")[:280] if latest_response else "None")
        follow_up_instruction = (
            "This is a follow-up/reflection turn.\n"
            "Acknowledge what the student got right in one sentence.\n"
            "Then add only missing points and one targeted next question.\n"
            "Do NOT restate the full definition. Keep under 220 words."
            if is_follow_up_turn
            else "If the current query overlaps with recent context, avoid repeating previous wording."
        )
        response_format = (
            "Response format:\n"
            "1) One validation sentence\n"
            "2) Incremental explanation (3-6 short lines)\n"
            "3) One guiding question"
            if is_follow_up_turn
            else (
                "Response format:\n"
                "1) Direct answer (2-8 short paragraphs)\n"
                "2) If useful, bullet list of steps\n"
                "3) Two guiding questions\n"
                "4) Never label the topic as 'general' if the query clearly names a topic."
            )
        )
        user_prompt = (
            f"Student query: {query}\n"
            f"Target concept: {concept}\n"
            f"Style: {style}\n"
            f"Style guidance: {style_instruction}\n"
            f"Student context: {context}\n"
            f"Latest student query in this thread: {previous_query}\n"
            f"Latest mentor answer summary: {previous_answer_summary}\n"
            f"Follow-up mode: {'yes' if is_follow_up_turn else 'no'}\n"
            f"Instruction: {follow_up_instruction}\n"
            f"Recent interaction context:\n{self._recent_context(student_id=student_id)}\n\n"
            f"{response_format}"
        )

        headers = None
        if api_base and "openrouter" in api_base.lower():
            headers = {
                "HTTP-Referer": os.getenv("OPENROUTER_SITE_URL", "http://localhost"),
                "X-Title": os.getenv("OPENROUTER_APP_NAME", "Mentor AI"),
            }

        errors: List[str] = []
        try:
            client = get_openai_client()
            for model in candidate_models:
                try:
                    completion = client.chat.completions.create(
                        model=model,
                        messages=[
                            {"role": "system", "content": system_prompt},
                            {"role": "user", "content": user_prompt},
                        ],
                        temperature=0.4,
                        max_tokens=700,
                        extra_headers=headers,
                        timeout=8,
                    )
                    content = completion.choices[0].message.content if completion.choices else None
                    if content:
                        MentorAIService._llm_backoff_until = 0.0
                        return content.strip()
                except Exception as model_error:
                    model_error_text = str(model_error)
                    errors.append(f"{model}: {model_error_text}")
                    lower_error = model_error_text.lower()
                    if any(token in lower_error for token in ["connection error", "timed out", "unauthorized", "invalid api key"]):
                        MentorAIService._llm_backoff_until = time.time() + 180
                        break
        except Exception as e:
            logger.warning("LLM client init failed; falling back to local templates. Error=%s", str(e))
            return None

        if errors:
            # Keep it compact (avoid spewing many lines into dev terminals).
            compact = "; ".join(errors[:5])
            suffix = " ..." if len(errors) > 5 else ""
            logger.warning("LLM response failed; falling back to local templates. Errors=%s%s", compact, suffix)

        return None

    @staticmethod
    def _concept_definition(concept: str) -> str:
        concept_l = (concept or "general").lower()
        definitions = {
            "machine learning": (
                "Machine learning is a method of building models that learn patterns from data "
                "to make predictions or decisions without writing explicit rules for every case."
            ),
            "gradient descent": (
                "Gradient descent is an optimization algorithm that updates parameters by moving "
                "in the negative gradient direction to reduce a loss function."
            ),
            "backpropagation": (
                "Backpropagation is the chain-rule based process used in neural networks to compute "
                "how each weight affects final loss, enabling targeted weight updates."
            ),
            "data analysis": (
                "Data analysis is the process of cleaning, exploring, and interpreting data "
                "to answer questions and support decisions."
            ),
            "data engineering": (
                "Data engineering designs and operates data pipelines that ingest, transform, and load data "
                "into reliable systems for analytics and machine learning."
            ),
            "sql": (
                "SQL is the standard language for querying and transforming structured data "
                "stored in relational databases."
            ),
            "python": (
                "Python is a high-level programming language widely used for automation, "
                "data analysis, machine learning, and backend development."
            ),
            "statistics": (
                "Statistics is the science of collecting, summarizing, and reasoning about data "
                "to make decisions under uncertainty."
            ),
            "linear algebra": (
                "Linear algebra studies vectors, matrices, and linear transformations; "
                "it is foundational for machine learning and optimization."
            ),
            "neural networks": (
                "Neural networks are layered function approximators that transform inputs through "
                "learned weights and nonlinear activations."
            ),
            "deep learning": (
                "Deep learning is a subset of machine learning that uses multi-layer neural networks "
                "to learn complex patterns from large datasets."
            ),
            "supervised learning": (
                "Supervised learning trains a model using labeled examples (input-output pairs) "
                "to predict known target variables."
            ),
            "unsupervised learning": (
                "Unsupervised learning finds structure or patterns in unlabeled data, "
                "such as clusters, latent factors, or anomalies."
            ),
            "overfitting": (
                "Overfitting happens when a model learns noise and dataset-specific quirks, "
                "performing well on training data but poorly on unseen data."
            ),
            "underfitting": (
                "Underfitting happens when a model is too simple to capture real signal, "
                "so it performs poorly on both training and test data."
            ),
            "regularization": (
                "Regularization controls model complexity to improve generalization, "
                "for example by penalizing large weights (L1/L2) or using dropout/early stopping."
            ),
            "system design": (
                "System design is the process of architecting scalable, reliable, and maintainable software systems "
                "by defining components, data flow, and trade-offs."
            ),
            "data structures": (
                "Data structures are ways to organize and store data efficiently, enabling faster operations "
                "like search, insert, delete, and traversal."
            ),
            "algorithms": (
                "Algorithms are step-by-step procedures to solve computational problems with measurable "
                "time and space complexity."
            ),
            "api design": (
                "API design defines how systems communicate through clear endpoints, request/response contracts, "
                "versioning, and error handling."
            ),
            "kubernetes": (
                "Kubernetes is a container orchestration platform that automates deployment, scaling, and management "
                "of containerized applications."
            ),
        }
        return definitions.get(
            concept_l,
            f"{concept} is a topic where you model the core idea, apply it on examples, and validate outcomes with measurable criteria.",
        )

    @staticmethod
    def _comparison_response(query_l: str) -> Optional[str]:
        pattern = r"(?:difference between|compare)\s+([a-z0-9\-\s]+?)\s+(?:and|vs)\s+([a-z0-9\-\s]+)"
        match = re.search(pattern, query_l)
        if not match:
            return None

        left = " ".join(match.group(1).split())
        right = " ".join(match.group(2).split())

        if "supervised" in left and "unsupervised" in right or "supervised" in right and "unsupervised" in left:
            return (
                "Supervised vs Unsupervised Learning:\n"
                "- Data: supervised uses labeled data; unsupervised uses unlabeled data.\n"
                "- Goal: supervised predicts targets; unsupervised discovers hidden structure.\n"
                "- Typical tasks: supervised -> classification/regression, unsupervised -> clustering/dimensionality reduction.\n"
                "- Evaluation: supervised has clear metrics against labels; unsupervised often needs proxy/business evaluation.\n"
                "- Use supervised when you have reliable labels; use unsupervised when labels are unavailable and exploration is needed."
            )

        return (
            f"Comparison: {left} vs {right}\n"
            f"- Purpose: What problem does each solve?\n"
            f"- Input/Output: How do inputs and outputs differ?\n"
            f"- Strengths: Where is {left} better, and where is {right} better?\n"
            f"- Trade-offs: Accuracy, speed, data requirement, interpretability.\n\n"
            f"To make this concrete, tell me your use case and I will recommend one with reasons."
        )

    @staticmethod
    def _local_socratic_response(
        query: str,
        concept: str,
        style: str,
        latest_response: Optional[MentorResponse] = None,
        is_follow_up_turn: bool = False,
    ) -> str:
        query_l = (query or "").lower()
        extracted_topic = MentorAIService._extract_topic_from_query(query)
        topic = concept if concept and concept != "general" else extracted_topic
        if not topic or topic == "general":
            topic = "the topic in your question"

        if is_follow_up_turn and latest_response:
            previous_concept = concept if concept and concept != "general" else (latest_response.target_concept or topic)
            if previous_concept == "data engineering":
                return (
                    "Good start. You correctly identified ETL in data engineering.\n"
                    "To strengthen your answer, add what happens inside each step:\n"
                    "1. Extract: source reliability and ingestion frequency.\n"
                    "2. Transform: cleaning, type normalization, deduplication, quality rules.\n"
                    "3. Load: warehouse/lake target schema and validation checks.\n\n"
                    "Now apply it: if clickstream data has duplicate events, which transformation rule would you use before load?"
                )

            return (
                f"Good follow-up. Your summary of {previous_concept} is on track.\n"
                "To make it complete, add:\n"
                "1. Why each step exists.\n"
                "2. One concrete tool/example.\n"
                "3. One common failure mode and prevention check.\n\n"
                f"Can you extend your answer with one practical {previous_concept} example?"
            )
        asks_definition = (
            query_l.startswith("what is ")
            or query_l.startswith("what are ")
            or query_l.startswith("define ")
            or " meaning of " in query_l
        )
        asks_why = query_l.startswith("why ") or " why " in query_l
        asks_how = query_l.startswith("how ") or " how " in query_l
        asks_compare = "difference between" in query_l or " compare " in query_l or " vs " in query_l
        asks_for_roadmap = any(token in query_l for token in ["roadmap", "road map", "plan", "path", "career", "become", "from scratch"])
        roadmap_continuation = any(
            token in query_l
            for token in ["current level", "timeline", "month", "months", "week", "weeks", "beginner", "intermediate", "advanced"]
        )
        asks_data_analyst = any(token in query_l for token in ["data analyst", "data analysis", "analytics"])
        asks_ml = any(token in query_l for token in ["machine learning", "ml"])
        if concept == "machine learning":
            asks_ml = True

        if asks_compare:
            compare = MentorAIService._comparison_response(query_l)
            if compare:
                return (
                    f"{compare}\n\n"
                    "Guiding questions:\n"
                    "1. What is your dataset size/type?\n"
                    "2. Is interpretability or raw accuracy more important?"
                )

        if (asks_for_roadmap or roadmap_continuation) and asks_data_analyst:
            return (
                "Great goal. Here is a practical 16-week roadmap to become a data analyst:\n\n"
                "Phase 1 (Weeks 1-4): Foundations\n"
                "- Excel: lookup functions, pivot tables, data cleaning\n"
                "- SQL basics: SELECT, WHERE, GROUP BY, JOIN\n"
                "- Statistics: mean/median, variance, distributions, hypothesis basics\n"
                "Output: 2 mini tasks solved in Excel + SQL\n\n"
                "Phase 2 (Weeks 5-8): Python for analysis\n"
                "- Python + pandas + numpy for cleaning and exploration\n"
                "- Matplotlib/Seaborn for visualization\n"
                "- EDA workflow: missing values, outliers, feature summaries\n"
                "Output: 2 EDA notebooks on public datasets\n\n"
                "Phase 3 (Weeks 9-12): BI + dashboards\n"
                "- Power BI or Tableau (pick one)\n"
                "- Build KPI dashboard with filters, drill-downs, and story flow\n"
                "- Learn metric design: conversion, retention, cohort views\n"
                "Output: 2 portfolio dashboards\n\n"
                "Phase 4 (Weeks 13-16): Portfolio + interview prep\n"
                "- 3 end-to-end projects (business problem -> analysis -> recommendation)\n"
                "- SQL interview practice: window functions, case expressions, joins\n"
                "- Communicate findings in 1-page business summary per project\n"
                "Output: GitHub portfolio + resume bullets + mock interview set\n\n"
                "Weekly routine (recommended):\n"
                "- 5 days x 2 hours learning, 1 day project building, 1 day revision\n\n"
                "Next Socratic step:\n"
                "1. Which tools do you already know today (Excel/SQL/Python/BI)?\n"
                "2. Can you commit to a realistic weekly hour target?"
            )

        if (asks_for_roadmap or roadmap_continuation) and asks_ml:
            return (
                "Great target. Here is a practical 5-month roadmap for machine learning (beginner-friendly):\n\n"
                "Month 1: Math + Python foundations\n"
                "- Python basics, numpy, pandas\n"
                "- Core math: linear algebra basics, derivatives, probability essentials\n"
                "Output: 2 notebooks (EDA + math intuition)\n\n"
                "Month 2: Classical ML core\n"
                "- Regression, classification, metrics (accuracy, precision/recall, ROC-AUC)\n"
                "- Train/validation/test split, cross-validation, bias-variance\n"
                "Output: 3 classical ML mini-projects\n\n"
                "Month 3: Feature engineering + model improvement\n"
                "- Missing values, encoding, scaling, leakage checks\n"
                "- Hyperparameter tuning (Grid/Random Search)\n"
                "Output: one polished end-to-end tabular ML project\n\n"
                "Month 4: Intro deep learning\n"
                "- Neural network basics, backprop intuition, optimization basics\n"
                "- Build simple models in TensorFlow/PyTorch\n"
                "Output: 1 DL project (image or text starter)\n\n"
                "Month 5: Portfolio + interview prep\n"
                "- 2 strong projects with clear problem, metrics, and business impact\n"
                "- SQL + ML interview prep and model explainability (SHAP/feature importance)\n"
                "Output: GitHub portfolio + resume bullets + mock interview answers\n\n"
                "Weekly plan:\n"
                "- 10-12 hours: 6 hours learning, 4 hours projects, 1-2 hours revision\n\n"
                "To personalize this to your level:\n"
                "1. Are you stronger in coding or math right now?\n"
                "2. Do you want job-ready Data Analyst -> ML path, or direct ML Engineer path?"
            )

        if asks_definition:
            normalized_topic = MentorAIService._normalize_topic_alias(topic)
            definition = MentorAIService._concept_definition(normalized_topic)
            return (
                f"{definition}\n\n"
                "Quick example:\n"
                f"- In {normalized_topic}, you define an objective, learn from data/examples, and evaluate with clear metrics.\n\n"
                "Check your understanding:\n"
                "1. Can you explain this in one sentence to a beginner?\n"
                "2. What is one real-world use case you care about?"
            )

        if "overfitting" in query_l:
            return (
                "Overfitting happens when model complexity is high relative to data signal, so the model memorizes training noise.\n"
                "Signs: very low training error but much higher validation/test error.\n"
                "Common causes: too many parameters, small dataset, weak regularization, data leakage.\n"
                "Fixes: cross-validation, simpler model, more data/augmentation, stronger regularization, early stopping.\n\n"
                "Guiding questions:\n"
                "1. What is your train vs validation performance gap?\n"
                "2. Which anti-overfitting method have you already tried?"
            )

        if "underfitting" in query_l:
            return (
                "Underfitting means the model is too simple or not trained enough to capture signal.\n"
                "Signs: both training and validation errors remain high.\n"
                "Fixes: richer features, higher-capacity model, better optimization, longer training, reduced regularization.\n\n"
                "Guiding questions:\n"
                "1. Is training error already high?\n"
                "2. Which part of your pipeline limits model capacity most?"
            )

        if "hashing" in query_l or "hash table" in query_l:
            return (
                "We use hashing because it gives near O(1) average-time lookup, insert, and delete by mapping keys to array indices.\n"
                "That makes dictionaries/maps fast for membership checks and key-value access.\n"
                "Trade-offs: collisions can degrade performance, and hash tables are usually unordered unless explicitly designed otherwise.\n\n"
                "Guiding questions:\n"
                "1. What collision strategy are you using (chaining or open addressing)?\n"
                "2. How does load factor affect performance in your implementation?"
            )

        if asks_for_roadmap or roadmap_continuation:
            return (
                f"Here is a practical roadmap for {topic}:\n"
                "Phase 1 (Weeks 1-2): Fundamentals\n"
                "- Learn core terminology and key principles\n"
                "- Solve 5 beginner exercises\n\n"
                "Phase 2 (Weeks 3-6): Applied practice\n"
                "- Work through guided tutorials and mini-projects\n"
                "- Build one end-to-end project you can explain\n\n"
                "Phase 3 (Weeks 7-10): Intermediate depth\n"
                "- Study common mistakes, edge cases, and trade-offs\n"
                "- Build a second project with better structure and documentation\n\n"
                "Phase 4 (Weeks 11-12): Portfolio and revision\n"
                "- Prepare summary notes and interview-style Q&A\n"
                "- Publish projects and track measurable improvement\n\n"
                "To personalize this roadmap:\n"
                "1. What is your current level?\n"
                "2. What timeline are you targeting?"
            )

        if asks_why:
            return (
                f"Great question. For '{query.strip()}', the core reason usually comes from how {topic} balances objective, constraints, and assumptions.\n"
                "A practical way to reason about it:\n"
                "1) Identify the goal,\n"
                "2) Identify what changes the result most,\n"
                "3) Check one case where the rule fails.\n\n"
                "Guiding questions:\n"
                "1. What assumption are you making implicitly?\n"
                "2. What breaks if that assumption fails?"
            )

        if asks_how:
            return (
                f"How to approach {topic} (practical view):\n"
                "- Step 1: define the objective and success metric\n"
                "- Step 2: break the task into smaller steps\n"
                "- Step 3: apply the method with one concrete example\n"
                "- Step 4: review result and improve the weak step\n\n"
                "Guiding questions:\n"
                "1. Which step is unclear for you right now?\n"
                "2. Want me to walk through a small worked example?"
            )

        # Concept-aware responses for common core questions so fallback mode
        # still produces useful mentoring output when no LLM key is configured.
        if concept == "gradient descent" and ("negative gradient" in query_l or "why" in query_l):
            if style == "simple":
                return (
                    "In gradient descent, the gradient points in the direction where loss increases fastest.\n"
                    "So we move in the *negative* gradient direction to make the loss go down.\n"
                    "Update rule: theta <- theta - alpha * grad(J(theta)).\n\n"
                    "Check your understanding:\n"
                    "1. If gradient is positive, should theta increase or decrease?\n"
                    "2. What happens if alpha (learning rate) is too large?\n"
                    "3. Can you picture this as walking downhill on a hill?"
                )

            if style == "deep":
                return (
                    "For differentiable J(theta), grad J(theta) gives steepest ascent under Euclidean norm.\n"
                    "Hence the steepest descent direction is -grad J(theta). A first-order step gives:\n"
                    "J(theta + d) ~ J(theta) + grad J(theta)^T d.\n"
                    "Choosing d = -alpha grad J(theta) yields\n"
                    "J(theta + d) ~ J(theta) - alpha ||grad J(theta)||^2, which decreases for alpha > 0 (locally).\n\n"
                    "Probe this further:\n"
                    "1. Which assumptions make this local decrease argument valid?\n"
                    "2. How does non-convexity change the guarantee?\n"
                    "3. Why do we often normalize or adapt the step size?"
                )

            return (
                "Great question. We use the negative gradient because the gradient points to fastest *increase* in loss,\n"
                "and training needs fastest *decrease*. So each step is:\n"
                "theta_new = theta_old - alpha * grad(J(theta)).\n"
                "Intuitively: gradient is an uphill arrow; putting a minus sign makes you walk downhill.\n\n"
                "Think through these:\n"
                "1. If grad(J)=0, what does that say about the point?\n"
                "2. How does alpha control stability vs speed?\n"
                "3. Why might momentum help compared with plain gradient descent?"
            )

        if concept == "backpropagation" and ("why" in query_l or "how" in query_l):
            return (
                "Backpropagation computes how each weight affects final loss using the chain rule,\n"
                "so updates are not random; they are targeted to reduce error.\n"
                "At each layer, we pass gradients backward and apply:\n"
                "w <- w - alpha * (dL/dw).\n\n"
                "Use this checklist:\n"
                "1. Where does the chain rule appear in one hidden-layer network?\n"
                "2. Why do activation derivatives matter for gradient flow?\n"
                "3. What failure do you get when gradients vanish?"
            )

        if style == "simple":
            return (
                f"Here's a simple explanation of {topic}:\n"
                f"- Start with the core idea behind {topic}.\n"
                "- Work one tiny example end-to-end.\n"
                "- Verify what changes the output most.\n\n"
                "Now reflect:\n"
                "1. Which single step is still unclear?\n"
                "2. Can you restate the idea in one sentence?\n"
                "3. What example would you try next?"
            )

        if style == "deep":
            return (
                f"For a deeper look at {topic}, start from definitions and derive the core relation.\n"
                "Then test assumptions and failure modes, not just the happy path.\n\n"
                "Challenge questions:\n"
                "1. Which theorem or property justifies each step?\n"
                "2. What edge case breaks the derivation?\n"
                "3. How would you modify the method for that edge case?"
            )

        return (
            f"Great question on {topic}. Here is a practical explanation:\n"
            f"- Objective: what {topic} is trying to optimize or explain.\n"
            "- Mechanism: the key step-by-step transformation.\n"
            "- Outcome: how you verify the method worked.\n\n"
            "Use these to self-check:\n"
            "1. What is the objective of this method?\n"
            "2. Which inputs control the output most strongly?\n"
            "3. Can you solve one small example without skipping steps?"
        )

    @staticmethod
    def _generate_guiding_question(concept: str, style: str) -> str:
        bank = {
            "simple": [
                f"What is the first principle behind {concept}?",
                f"Can you give a one-line definition of {concept} in your own words?",
            ],
            "conceptual": [
                f"How does {concept} connect to something you already understand?",
                f"Which variable in {concept} changes the outcome most?",
            ],
            "deep": [
                f"What assumptions make the standard derivation of {concept} valid?",
                f"How would you formally justify each step in {concept}?",
            ],
        }
        return random.choice(bank.get(style, bank["conceptual"]))

    def _store_response(
        self,
        student_id: int,
        query: str,
        response: str,
        style: str,
        concept: str,
        confidence: float,
    ) -> str:
        response_id = str(uuid.uuid4())

        snapshot_rows = (
            self.db.query(WeaknessScore)
            .filter(WeaknessScore.student_id == student_id)
            .order_by(desc(WeaknessScore.weakness_score))
            .limit(5)
            .all()
        )
        weakness_snapshot = {row.concept_name: round(row.weakness_score, 3) for row in snapshot_rows}

        entity = MentorResponse(
            response_id=response_id,
            student_id=student_id,
            student_weakness_state=weakness_snapshot,
            student_confidence=confidence,
            query=query,
            response=response,
            explanation_style=style,
            target_concept=concept,
        )
        self.db.add(entity)
        self.db.commit()
        return response_id


class FeedbackService:
    """Human-in-the-loop feedback processor."""

    def __init__(self, db: Session):
        self.db = db
        self.profile_service = StudentProfileService(db)
        self.weakness_service = WeaknessAnalyzerService(db)

    def submit_feedback(
        self,
        student_id: int,
        response_id: str,
        feedback_type: object,
        rating: Optional[float] = None,
        comments: Optional[str] = None,
        focus_concept: Optional[str] = None,
    ) -> Tuple[Feedback, Optional[AdaptationUpdate]]:
        profile = self.profile_service.get_profile(student_id)
        if not profile:
            raise ValueError(f"Profile not found for student {student_id}")

        feedback_enum = _coerce_feedback_type(feedback_type)

        feedback = Feedback(
            student_id=student_id,
            response_id=response_id,
            feedback_type=feedback_enum,
            rating=rating,
            comments=comments,
            focus_concept=(focus_concept or None),
        )
        self.db.add(feedback)
        self.db.commit()
        self.db.refresh(feedback)

        adaptation = self._adapt_to_feedback(student_id, feedback_enum, rating)
        self._apply_feedback_to_weakness(student_id, feedback_enum, focus_concept)

        return feedback, adaptation

    def _adapt_to_feedback(
        self,
        student_id: int,
        feedback_type: FeedbackType,
        rating: Optional[float] = None,
    ) -> Optional[AdaptationUpdate]:
        profile = self.profile_service.get_profile(student_id)
        if not profile:
            return None

        old_difficulty = profile.preferred_difficulty.value
        old_confidence = profile.confidence_level
        reasons: List[str] = []

        if feedback_type == FeedbackType.TOO_EASY:
            if profile.preferred_difficulty == DifficultyLevel.EASY:
                profile.preferred_difficulty = DifficultyLevel.MEDIUM
                reasons.append("Raised difficulty to medium after repeated easy feedback")
            elif profile.preferred_difficulty == DifficultyLevel.MEDIUM:
                profile.preferred_difficulty = DifficultyLevel.HARD
                reasons.append("Raised difficulty to hard after easy feedback")

        if feedback_type == FeedbackType.TOO_HARD:
            if profile.preferred_difficulty == DifficultyLevel.HARD:
                profile.preferred_difficulty = DifficultyLevel.MEDIUM
                reasons.append("Reduced difficulty to medium after hard feedback")
            elif profile.preferred_difficulty == DifficultyLevel.MEDIUM:
                profile.preferred_difficulty = DifficultyLevel.EASY
                reasons.append("Reduced difficulty to easy after hard feedback")

        if feedback_type == FeedbackType.UNCLEAR:
            profile.confidence_level = max(0.0, profile.confidence_level - 0.05)
            reasons.append("Reduced confidence slightly due to unclear feedback")

        if rating is not None:
            if rating <= 2.0:
                profile.confidence_level = max(0.0, profile.confidence_level - 0.1)
                reasons.append("Lowered confidence estimate from low rating")
            elif rating >= 4.0:
                profile.confidence_level = min(1.0, profile.confidence_level + 0.07)
                reasons.append("Raised confidence estimate from positive rating")

        profile.updated_at = datetime.utcnow()
        self.db.commit()

        new_difficulty = profile.preferred_difficulty.value
        if not reasons and new_difficulty == old_difficulty and profile.confidence_level == old_confidence:
            return None

        return AdaptationUpdate(
            previous_difficulty=old_difficulty,
            new_difficulty=new_difficulty,
            adjustment_reason="; ".join(reasons) if reasons else "No major adaptation required",
            confidence_change=round(profile.confidence_level, 3),
        )

    def _apply_feedback_to_weakness(
        self,
        student_id: int,
        feedback_type: FeedbackType,
        focus_concept: Optional[str],
    ) -> None:
        if not focus_concept:
            return

        weakness = self.weakness_service.get_or_create_weakness(student_id, focus_concept)

        delta_map = {
            FeedbackType.TOO_HARD: 0.08,
            FeedbackType.UNCLEAR: 0.05,
            FeedbackType.HELPFUL: -0.03,
            FeedbackType.TOO_EASY: -0.04,
        }
        delta = delta_map.get(feedback_type, 0.0)

        weakness.weakness_score = min(1.0, max(0.0, weakness.weakness_score + delta))
        weakness.last_updated = datetime.utcnow()
        self.db.commit()


class AdaptiveLearningService:
    """Adaptive learning orchestrator."""

    def __init__(self, db: Session):
        self.db = db
        self.profile_service = StudentProfileService(db)
        self.weakness_service = WeaknessAnalyzerService(db)

    def create_session(self, student_id: int, topic: str, difficulty_level: object) -> AdaptiveSession:
        profile = self.profile_service.get_profile(student_id)
        if not profile:
            raise ValueError(f"Profile not found for student {student_id}")

        snapshot = self.get_student_context_snapshot(student_id)
        session = AdaptiveSession(
            student_id=student_id,
            topic=(topic or "general"),
            difficulty_level=_coerce_difficulty(difficulty_level, profile.preferred_difficulty).value,
            interaction_count=0,
            context_snapshot=snapshot.model_dump(),
        )
        self.db.add(session)
        self.db.commit()
        self.db.refresh(session)
        return session

    def get_student_context_snapshot(self, student_id: int) -> StudentContextSnapshot:
        profile = self.profile_service.get_profile(student_id)
        if not profile:
            raise ValueError(f"Profile not found for student {student_id}")

        weaknesses = self.weakness_service.get_weakest_concepts(student_id, limit=5)
        strength_areas = self.weakness_service.get_strength_areas(student_id, limit=3)
        recent_feedback = (
            self.db.query(Feedback)
            .filter(Feedback.student_id == student_id)
            .order_by(desc(Feedback.created_at))
            .limit(10)
            .all()
        )

        return StudentContextSnapshot(
            confidence_level=round(profile.confidence_level, 3),
            primary_weakness_concepts=[row.concept_name for row in weaknesses],
            strength_areas=strength_areas,
            preferred_difficulty=profile.preferred_difficulty.value,
            recent_feedback_sentiment=self._analyze_feedback_sentiment(recent_feedback),
        )

    def generate_recommendations(self, student_id: int) -> List[Dict]:
        context = self.get_student_context_snapshot(student_id)
        profile = self.profile_service.get_profile(student_id)
        weaknesses = self.weakness_service.get_weakest_concepts(student_id, limit=3)
        recommendations: List[Dict] = []

        for row in weaknesses:
            if row.weakness_score >= 0.5:
                recommendations.append(
                    {
                        "priority": "high",
                        "recommendation_type": "Concept Review",
                        "suggested_action": f"Practice {row.concept_name} with 3 scaffolded questions",
                        "explanation": f"Weakness score is {row.weakness_score:.2f}, indicating repeated difficulty.",
                    }
                )

        if context.recent_feedback_sentiment == "negative":
            recommendations.append(
                {
                    "priority": "high",
                    "recommendation_type": "Difficulty Adjustment",
                    "suggested_action": "Use simpler explanations for next session and increase step-by-step prompts",
                    "explanation": "Recent feedback trend is negative (too_hard/unclear dominates).",
                }
            )

        if context.confidence_level < 0.45:
            focus = context.strength_areas[0] if context.strength_areas else "a familiar topic"
            recommendations.append(
                {
                    "priority": "medium",
                    "recommendation_type": "Confidence Boost",
                    "suggested_action": f"Start next session with a quick win in {focus}",
                    "explanation": "Confidence level is below 0.45; quick successes can improve learning momentum.",
                }
            )

        if profile and profile.preferred_difficulty == DifficultyLevel.HARD and context.recent_feedback_sentiment == "positive":
            recommendations.append(
                {
                    "priority": "low",
                    "recommendation_type": "Challenge Extension",
                    "suggested_action": "Add one transfer question connecting two concepts",
                    "explanation": "Positive feedback at high difficulty suggests readiness for integrative tasks.",
                }
            )

        if not recommendations:
            recommendations.append(
                {
                    "priority": "low",
                    "recommendation_type": "Steady Practice",
                    "suggested_action": "Continue current plan and review one recent concept",
                    "explanation": "No major risk signals detected from recent data.",
                }
            )

        return recommendations

    def generate_study_plan(
        self,
        student_id: int,
        weeks: int = 2,
        days_per_week: int = 5,
        daily_minutes: int = 60,
    ) -> Dict:
        profile = self.profile_service.get_profile(student_id)
        if not profile:
            raise ValueError(f"Profile not found for student {student_id}")

        weaknesses = self.weakness_service.get_weakest_concepts(
            student_id,
            limit=max(5, days_per_week),
        )
        prioritized_concepts = [row.concept_name for row in weaknesses if row.concept_name]
        if not prioritized_concepts:
            prioritized_concepts = self._fallback_plan_concepts(profile)

        key_weaknesses = [
            {
                "concept": row.concept_name,
                "weakness_score": round(row.weakness_score, 3),
                "priority": WeaknessAnalyzerService._calculate_learning_priority(row.weakness_score),
            }
            for row in weaknesses[:5]
        ]
        if not key_weaknesses:
            key_weaknesses = [
                {
                    "concept": concept,
                    "weakness_score": None,
                    "priority": "medium",
                }
                for concept in prioritized_concepts[:3]
            ]

        goal_text = (profile.goals or "").strip() or "Build stronger fundamentals with consistent deliberate practice."
        weekly_roadmap: List[Dict] = []
        for week_number in range(1, weeks + 1):
            primary_concept = prioritized_concepts[(week_number - 1) % len(prioritized_concepts)]
            secondary_concept = prioritized_concepts[week_number % len(prioritized_concepts)]

            days: List[Dict] = []
            for day_number in range(1, days_per_week + 1):
                concept_index = ((week_number - 1) * days_per_week + (day_number - 1)) % len(prioritized_concepts)
                focus_concept = prioritized_concepts[concept_index]

                days.append(
                    {
                        "day_number": day_number,
                        "focus_concept": focus_concept,
                        "objective": self._day_objective(day_number, focus_concept, primary_concept),
                        "activities": self._day_activities(
                            concept=focus_concept,
                            day_number=day_number,
                            daily_minutes=daily_minutes,
                            preferred_difficulty=profile.preferred_difficulty.value,
                        ),
                        "estimated_minutes": daily_minutes,
                    }
                )

            weekly_focus = (
                f"Strengthen {primary_concept} and connect it with {secondary_concept}."
                if primary_concept != secondary_concept
                else f"Deepen understanding and speed in {primary_concept}."
            )
            weekly_roadmap.append(
                {
                    "week_number": week_number,
                    "weekly_focus": weekly_focus,
                    "goal_alignment": self._goal_alignment(goal_text, primary_concept),
                    "days": days,
                }
            )

        return {
            "student_id": student_id,
            "goals": goal_text,
            "confidence_level": round(profile.confidence_level, 3),
            "preferred_difficulty": profile.preferred_difficulty.value,
            "weeks": weeks,
            "days_per_week": days_per_week,
            "daily_minutes": daily_minutes,
            "key_weaknesses": key_weaknesses,
            "weekly_roadmap": weekly_roadmap,
            "guidance": self._plan_guidance(
                confidence_level=profile.confidence_level,
                preferred_difficulty=profile.preferred_difficulty.value,
            ),
        }

    @staticmethod
    def _fallback_plan_concepts(profile: StudentProfile) -> List[str]:
        candidates: List[str] = []
        for item in (profile.skills or []) + (profile.interests or []):
            value = str(item).strip().lower()
            if value and value not in candidates:
                candidates.append(value)
            if len(candidates) >= 6:
                break

        if not candidates and profile.goals:
            for token in re.split(r"[,.;\n]", profile.goals.lower()):
                value = token.strip()
                if len(value) >= 3 and value not in candidates:
                    candidates.append(value)
                if len(candidates) >= 6:
                    break

        return candidates or ["core foundations", "problem solving", "revision"]

    @staticmethod
    def _day_objective(day_number: int, concept: str, weekly_anchor: str) -> str:
        cycle = {
            1: f"Rebuild fundamentals and identify one gap in {concept}.",
            2: f"Apply {concept} through guided practice with worked examples.",
            3: f"Strengthen recall speed and accuracy in {concept}.",
            4: f"Integrate {concept} with {weekly_anchor} in mixed problems.",
            0: f"Consolidate {concept} with reflection and a short self-quiz.",
        }
        return cycle[day_number % 5]

    @staticmethod
    def _day_activities(
        concept: str,
        day_number: int,
        daily_minutes: int,
        preferred_difficulty: str,
    ) -> List[str]:
        review_minutes = max(10, int(daily_minutes * 0.35))
        practice_minutes = max(15, int(daily_minutes * 0.45))
        reflection_minutes = max(5, daily_minutes - review_minutes - practice_minutes)

        problem_count = {
            "easy": "2-3",
            "medium": "3-4",
            "hard": "4-6",
        }.get(preferred_difficulty, "3-4")

        activities = [
            f"{review_minutes} min: review notes and one solved example for {concept}.",
            f"{practice_minutes} min: solve {problem_count} practice questions on {concept}.",
        ]

        if day_number % 5 == 0:
            activities.append(
                f"{reflection_minutes} min: take a short self-quiz and record the main mistakes to revisit."
            )
        else:
            activities.append(
                f"{reflection_minutes} min: summarize key takeaways and write one follow-up question."
            )

        return activities

    @staticmethod
    def _goal_alignment(goal_text: str, concept: str) -> str:
        compact_goal = " ".join(goal_text.split())
        if len(compact_goal) > 120:
            compact_goal = f"{compact_goal[:117]}..."
        return f"Focus on {concept} to move toward goal: {compact_goal}"

    @staticmethod
    def _plan_guidance(confidence_level: float, preferred_difficulty: str) -> List[str]:
        guidance = [
            "Track completion daily and carry unfinished tasks to the next day.",
            "After each week, re-check weakness scores and rotate concepts if needed.",
        ]

        if confidence_level < 0.45:
            guidance.append("Start each day with one easy warm-up question to build momentum.")
        if preferred_difficulty == "hard":
            guidance.append("End each study day with one transfer problem that mixes multiple concepts.")

        return guidance

    @staticmethod
    def _analyze_feedback_sentiment(feedbacks: List[Feedback]) -> str:
        if not feedbacks:
            return "neutral"

        score = 0
        for row in feedbacks:
            if row.feedback_type == FeedbackType.HELPFUL:
                score += 1
            elif row.feedback_type == FeedbackType.TOO_EASY:
                score += 0
            else:
                score -= 1

        if score > 0:
            return "positive"
        if score < 0:
            return "negative"
        return "neutral"


__all__ = [
    "StudentProfileService",
    "WeaknessAnalyzerService",
    "MentorAIService",
    "FeedbackService",
    "AdaptiveLearningService",
]
